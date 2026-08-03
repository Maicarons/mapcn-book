# -*- coding: utf-8 -*-
"""Build a teaching PPTX for 《Mapcn 从入门到进阶》 from the 14 chapter markdown files.

Mirrors the pipeline established in duckdb-1.5.5/mybook/tools/build_pptx.py,
adapted to the mapcn 4-part structure and monochrome design language.
"""
import os, re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- palette (monochrome + single restrained slate accent) ----------
PRIMARY   = RGBColor(0x18, 0x18, 0x1B)   # near-black charcoal
PRIMARY2  = RGBColor(0x3F, 0x3F, 0x46)   # zinc-700
ACCENT    = RGBColor(0x4F, 0x6B, 0x8A)   # muted slate-blue accent
LIGHT     = RGBColor(0xF4, 0xF4, 0xF5)   # zinc-100 panel
CODEBG    = RGBColor(0xEF, 0xEF, 0xF1)   # code box bg
DARK      = RGBColor(0x27, 0x27, 0x2B)   # near-black text
GRAY      = RGBColor(0x71, 0x71, 0x78)   # muted
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
CODETXT   = RGBColor(0x1F, 0x29, 0x33)
TITLEFONT = "Microsoft YaHei"
CODEFONT  = "Consolas"

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
CHAPTERS_DIR = os.path.join(REPO, "docs", "chapters")
OUT = os.path.join(REPO, "docs", "public", "downloads", "mapcn-teaching.pptx")

PARTS = [
    ("第一部分  入门与环境", [1, 2]),
    ("第二部分  地图基础与交互", [3, 4, 5, 6]),
    ("第三部分  图形与数据", [7, 8, 9, 10]),
    ("第四部分  进阶与实战", [11, 12, 13, 14]),
]

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
slide_no = 0

# ---------- helpers ----------
def set_run(run, size, bold=False, color=None, font=TITLEFONT, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {}); rPr.append(e)
        e.set('typeface', font)

def add_rect(slide, l, t, w, h, fill=None, line=None, line_w=None, shadow=False):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w or 1)
    return sp

def add_text(slide, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             wrap=True, space_after=4, line_spacing=1.0):
    """runs: list of paragraphs; each paragraph = list of (text, dict) tuples."""
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for text, opt in para:
            r = p.add_run(); r.text = text
            set_run(r, opt.get('size', 14), bold=opt.get('bold', False),
                    color=opt.get('color', DARK), font=opt.get('font', TITLEFONT),
                    italic=opt.get('italic', False))
    return tb

def style_bullet(paragraph, level=0, char='•'):
    pPr = paragraph._p.get_or_add_pPr()
    marL = Inches(0.32 + 0.30 * level)
    pPr.set('marL', str(int(marL)))
    pPr.set('indent', str(int(-Inches(0.18))))
    for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum'):
        e = pPr.find(qn(tag))
        if e is not None: pPr.remove(e)
    buFont = pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial', 'pitchFamily': '34', 'charset': '-1'})
    buChar = pPr.makeelement(qn('a:buChar'), {'char': char})
    pPr.append(buFont); pPr.append(buChar)

def add_bullets(slide, l, t, w, h, items, size=14, color=DARK, space=6, line_spacing=1.05):
    """items: list of (text, level)"""
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, (text, level) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space); p.space_before = Pt(0); p.line_spacing = line_spacing
        style_bullet(p, level=level)
        segs = re.split(r'(\*\*.+?\*\*)', text)
        for seg in segs:
            if not seg: continue
            if seg.startswith('**') and seg.endswith('**'):
                r = p.add_run(); r.text = seg[2:-2]; set_run(r, size, bold=True, color=PRIMARY if level == 0 else color)
            else:
                r = p.add_run(); r.text = seg; set_run(r, size, color=color, bold=(level == 0))
    return tb

def add_codebox(slide, l, t, w, h, code):
    box = add_rect(slide, l, t, w, h, fill=CODEBG, line=PRIMARY2, line_w=0.75)
    tb = slide.shapes.add_textbox(Inches(l + 0.12), Inches(t + 0.1), Inches(w - 0.24), Inches(h - 0.2))
    tf = tb.text_frame; tf.word_wrap = False
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    lines = code.split('\n')
    maxlines = max(1, int((h - 0.2) / 0.155))
    if len(lines) > maxlines:
        lines = lines[:maxlines]
        lines.append("…（完整代码与说明见教材对应章节）")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 0.95; p.space_after = Pt(0); p.space_before = Pt(0)
        r = p.add_run(); r.text = ln if ln else " "
        set_run(r, 9, color=CODETXT, font=CODEFONT)

def add_footer(slide, label):
    global slide_no
    slide_no += 1
    add_rect(slide, 0, 7.06, 13.333, 0.02, fill=PRIMARY)
    add_text(slide, 0.5, 7.1, 9.5, 0.3,
             [[("Mapcn 从入门到进阶  ·  教学讲义", {'size': 9, 'color': GRAY})]])
    add_text(slide, 11.0, 7.1, 1.83, 0.3,
             [[("%02d" % slide_no, {'size': 9, 'color': GRAY})]], align=PP_ALIGN.RIGHT)

def add_title_bar(slide, title):
    add_rect(slide, 0, 0, 13.333, 1.0, fill=PRIMARY)
    add_rect(slide, 0, 1.0, 13.333, 0.06, fill=ACCENT)
    add_text(slide, 0.5, 0, 12.3, 1.0,
             [[(title, {'size': 24, 'bold': True, 'color': WHITE})]], anchor=MSO_ANCHOR.MIDDLE)

def new_slide(bg=None):
    s = prs.slides.add_slide(BLANK)
    if bg is not None:
        s.background.fill.solid(); s.background.fill.fore_color.rgb = bg
    return s

def clean_title(title):
    """Strip a leading '第 X 章' chapter-number prefix for display."""
    return re.sub(r'^第\s*\d+\s*[.章]?\s*', '', title).strip()

# ---------- cover ----------
def build_cover():
    s = new_slide(bg=PRIMARY)
    add_rect(s, 0, 3.05, 13.333, 0.06, fill=ACCENT)
    add_text(s, 0.8, 1.9, 11.7, 1.2,
             [[("Mapcn 从入门到进阶", {'size': 46, 'bold': True, 'color': WHITE})],
              [("教学讲义", {'size': 46, 'bold': True, 'color': WHITE})]],
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 0.85, 3.25, 11.5, 0.6,
             [[("基于 mapcn 官方文档  ·  可复制粘贴的 React 地图组件库", {'size': 19, 'color': RGBColor(0xCF, 0xDE, 0xEF)})]])
    add_text(s, 0.85, 6.4, 11.5, 0.6,
             [[("面向 React 开发者的地图组件实战教程 —— 从基础接入到进阶封装", {'size': 13, 'color': RGBColor(0xB9, 0xCC, 0xE2)})],
              [("编制日期：2026-08-03", {'size': 11, 'color': GRAY})]])

# ---------- agenda ----------
def build_agenda(ch_titles):
    groups = [PARTS[0:2], PARTS[2:4]]
    for gi, grp in enumerate(groups):
        s = new_slide()
        add_title_bar(s, "课程大纲（%d/%d）" % (gi + 1, len(groups)))
        y = 1.4
        for pname, chs in grp:
            add_text(s, 0.5, y, 12.3, 0.38,
                     [[("%s" % pname, {'size': 15, 'bold': True, 'color': PRIMARY})]])
            y += 0.40
            inline = "　".join("第%d章 %s" % (c, clean_title(ch_titles[c])) for c in chs)
            add_text(s, 0.75, y, 12.0, 0.5,
                     [[(inline, {'size': 12.5, 'color': DARK})]], line_spacing=1.05)
            y += 0.62
        add_footer(s, "大纲")

# ---------- part divider ----------
def build_part_divider(pname, chs, ch_titles):
    s = new_slide(bg=PRIMARY)
    idx = [p[0] for p in PARTS].index(pname) + 1
    add_text(s, 0.9, 1.6, 11.5, 1.2,
             [[("PART %02d" % idx, {'size': 30, 'bold': True, 'color': ACCENT})]])
    add_rect(s, 0.95, 2.9, 2.2, 0.06, fill=ACCENT)
    add_text(s, 0.9, 3.0, 11.5, 1.0,
             [[(pname.split('  ', 1)[-1] if '  ' in pname else pname, {'size': 34, 'bold': True, 'color': WHITE})]],
             anchor=MSO_ANCHOR.MIDDLE)
    items = [("第 %d 章  %s" % (c, clean_title(ch_titles[c])), 0) for c in chs]
    add_bullets(s, 1.0, 4.0, 11.3, 0.42 * len(items) + 0.2, items, size=13, color=RGBColor(0xD7, 0xE3, 0xF2), space=3)
    add_footer(s, pname)

# ---------- chapter title slide ----------
def build_chapter_title(num, title, intro):
    s = new_slide(bg=LIGHT)
    add_rect(s, 0, 0, 0.28, 7.5, fill=PRIMARY)
    add_rect(s, 0.28, 0, 0.10, 7.5, fill=ACCENT)
    add_text(s, 0.9, 1.5, 11.4, 1.4,
             [[("第 %d 章" % num, {'size': 30, 'bold': True, 'color': ACCENT})]], anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, 0.9, 2.7, 11.4, 1.6,
             [[(clean_title(title), {'size': 34, 'bold': True, 'color': PRIMARY})]])
    intro_clean = re.sub(r'\s+', ' ', intro).strip()
    if len(intro_clean) > 130:
        intro_clean = intro_clean[:130] + "…"
    add_text(s, 0.95, 4.5, 11.3, 2.4,
             [[(intro_clean, {'size': 15, 'color': DARK})]], line_spacing=1.25)
    add_footer(s, "第 %d 章" % num)

# ---------- content slide ----------
def render_section(num, sec, ch_label):
    s = new_slide()
    add_title_bar(s, sec['title'])
    paras = sec['paras']
    subs  = sec['subs']
    codes = sec['codes']
    tables = sec['tables']
    has_code = len(codes) > 0
    has_tbl  = len(tables) > 0

    bullets = []
    for kind, txt in paras:
        t = txt.strip()
        if not t: continue
        t = re.sub(r'\s+', ' ', t)
        if len(t) > 90: t = t[:90] + "…"
        if kind == 'bullet':
            bullets.append((t, 1))
        else:
            bullets.append((t, 0))
    for sub in subs:
        st = re.sub(r'\s+', ' ', sub['title']).strip()
        bullets.append(("**%s**" % st, 0))
        if sub['text'].strip():
            tt = re.sub(r'\s+', ' ', sub['text']).strip()
            if len(tt) > 90: tt = tt[:90] + "…"
            bullets.append((tt, 1))
    if len(bullets) > 9:
        bullets = bullets[:9]
        bullets.append(("（更多要点与完整示例见教材对应章节）", 1))

    if has_code:
        add_bullets(s, 0.5, 1.3, 6.4, 5.5, bullets, size=13, space=7, line_spacing=1.05)
        add_codebox(s, 7.1, 1.3, 5.75, 5.5, codes[0])
    elif has_tbl:
        top = bullets[:3]
        add_bullets(s, 0.5, 1.3, 12.3, 1.0 + 0.42 * len(top), top, size=13, space=5, line_spacing=1.05)
        draw_table(s, 0.5, 1.45 + 0.42 * len(top), 12.3, tables[0])
    else:
        add_bullets(s, 0.5, 1.35, 12.3, 5.5, bullets, size=14.5, space=9, line_spacing=1.1)
    add_footer(s, ch_label)

def draw_table(slide, l, t, w, table):
    if not table: return
    max_bottom = 6.9
    avail = max_bottom - t
    nrows = min(len(table), max(3, int(avail / 0.34)))
    nrows = min(nrows, len(table))
    ncols = max(len(r) for r in table[:nrows])
    norm = [r + [''] * (ncols - len(r)) for r in table[:nrows]]
    colw = w / ncols
    rowh = min(0.5, avail / nrows)
    if rowh < 0.28: rowh = 0.28
    for c in range(ncols):
        add_rect(slide, l + c * colw, t, colw, rowh, fill=PRIMARY)
        add_text(slide, l + c * colw + 0.05, t, colw - 0.1, rowh,
                 [[(norm[0][c], {'size': 11, 'bold': True, 'color': WHITE})]], anchor=MSO_ANCHOR.MIDDLE)
    for r in range(1, nrows):
        y = t + r * rowh
        fill = LIGHT if r % 2 == 0 else WHITE
        for c in range(ncols):
            add_rect(slide, l + c * colw, y, colw, rowh, fill=fill, line=RGBColor(0xD0, 0xD8, 0xE2), line_w=0.5)
            add_text(slide, l + c * colw + 0.05, y, colw - 0.1, rowh,
                     [[(norm[r][c], {'size': 10.5, 'color': DARK})]], anchor=MSO_ANCHOR.MIDDLE)

# ---------- parsing ----------
HEAD_RE = re.compile(r'^#{1,6}\s*(第\s*\d+\s*章|\d+\.)')

def parse_blocks(lines):
    blocks = []
    i, n = 0, len(lines)
    cur_code = None; code_buf = []
    while i < n:
        line = lines[i]
        if line.strip().startswith('```'):
            if cur_code is None:
                cur_code = line.strip()[3:].strip(); code_buf = []
            else:
                blocks.append(('code', cur_code, '\n'.join(code_buf))); cur_code = None
            i += 1; continue
        if cur_code is not None:
            code_buf.append(line); i += 1; continue
        if '|' in line and i + 1 < n and re.match(r'^\s*\|?[\s:|-]+\|?\s*$', lines[i + 1]) and '-' in lines[i + 1]:
            tbl = []
            while i < n and '|' in lines[i] and lines[i].strip():
                row = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                tbl.append(row); i += 1
            if len(tbl) >= 2 and set(''.join(tbl[1])) <= set('-: |'):
                tbl = tbl[0:1] + tbl[2:]
            if tbl: blocks.append(('table', tbl))
            continue
        if line.startswith('#'):
            if HEAD_RE.match(line):
                level = len(line) - len(line.lstrip('#'))
                text = line.lstrip('#').strip()
                blocks.append(('h', level, text))
            i += 1; continue
        if line.strip().startswith('-') or line.strip().startswith('*'):
            blocks.append(('bullet', line.strip().lstrip('-*').strip())); i += 1; continue
        if line.strip() == '':
            i += 1; continue
        blocks.append(('para', line.strip())); i += 1
    return blocks

def parse_chapter(path, num):
    lines = open(path, encoding='utf-8').read().splitlines()
    blocks = parse_blocks(lines)
    title = None
    heading_levels = []
    for b in blocks:
        if b[0] == 'h':
            if title is None:
                title = b[2]
            else:
                heading_levels.append(b[1])
    if any(lv == 2 for lv in heading_levels):
        div = 2
    elif any(lv == 3 for lv in heading_levels):
        div = 3
    else:
        div = 2
    sections = []
    cur = None
    ch_intro_paras = []
    title_set = False
    for b in blocks:
        if b[0] == 'h':
            _, level, text = b
            if not title_set:
                title_set = True
                title = re.sub(r'^\d+\.\s*', '', text) if re.match(r'^\d+\.\s', text) else text
                continue
            if level <= div:
                cur = {'title': text, 'paras': [], 'subs': [], 'codes': [], 'tables': []}
                sections.append(cur)
            else:
                if cur is not None:
                    cur['subs'].append({'title': text, 'text': ''})
                else:
                    ch_intro_paras.append(('bullet', text))
        elif b[0] == 'para':
            if cur is None:
                ch_intro_paras.append(('para', b[1]))
            else:
                cur['paras'].append(('para', b[1]))
        elif b[0] == 'bullet':
            if cur is None:
                ch_intro_paras.append(('bullet', b[1]))
            else:
                cur['paras'].append(('bullet', b[1]))
        elif b[0] == 'code':
            if cur is not None:
                cur['codes'].append(b[2])
        elif b[0] == 'table':
            if cur is not None:
                cur['tables'].append(b[1])
    intro_txt = ' '.join(t for k, t in ch_intro_paras if k == 'para')
    if not intro_txt and ch_intro_paras:
        intro_txt = ch_intro_paras[0][1]
    sections = [s for s in sections if (s['paras'] or s['subs'] or s['codes'] or s['tables'])]
    return title, intro_txt, sections

# ---------- geometry QA ----------
def geometry_qa(tol=9525):
    W, H = prs.slide_width, prs.slide_height
    issues = 0
    for si, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            try:
                l, t, w, h = sh.left, sh.top, sh.width, sh.height
            except Exception:
                continue
            if l is None or t is None or w is None or h is None:
                continue
            if l < -tol or t < -tol or (l + w) > W + tol or (t + h) > H + tol:
                issues += 1
                print("  [QA] slide %02d overflow: L=%.3f T=%.3f R=%.3f B=%.3f" % (
                    si + 1, l / 914400.0, t / 914400.0,
                    (l + w) / 914400.0, (t + h) / 914400.0))
    return issues

# ---------- main ----------
def main():
    ch_titles = {}
    ch_data = {}
    for num in range(1, 15):
        path = os.path.join(CHAPTERS_DIR, "chapter%02d.md" % num)
        title, intro, sections = parse_chapter(path, num)
        ch_titles[num] = title
        ch_data[num] = (intro, sections)

    build_cover()
    build_agenda(ch_titles)

    for pname, chs in PARTS:
        build_part_divider(pname, chs, ch_titles)
        for c in chs:
            title = ch_titles[c]
            intro, sections = ch_data[c]
            build_chapter_title(c, title, intro)
            for sec in sections:
                render_section(c, sec, "第 %d 章" % c)

    issues = geometry_qa()
    prs.save(OUT)
    print("SAVED", OUT)
    print("total slides:", len(prs.slides._sldIdLst))
    print("geometry issues:", issues)

if __name__ == "__main__":
    main()
